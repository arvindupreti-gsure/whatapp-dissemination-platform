# -*- coding: utf-8 -*-
"""Report generation in Excel, CSV and PDF (FR-28 to FR-33)."""
import csv
import io
import datetime as dt

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import mm
from reportlab.platypus import (SimpleDocTemplate, Table, TableStyle, Paragraph,
                                Spacer)

import analytics
from db import SessionLocal, Campaign, Delivery, Group
from sqlalchemy import select

NAVY = "0E2A47"
LIGHT = "F1F5FA"


# ------------------------------------------------------------------- CSV
def to_csv(headers: list[str], rows: list[list]) -> bytes:
    buf = io.StringIO(newline="")
    w = csv.writer(buf, quoting=csv.QUOTE_ALL)
    w.writerow(headers)
    w.writerows(rows)
    return buf.getvalue().encode("utf-8-sig")


# ------------------------------------------------------------------ Excel
def to_xlsx(sheets: dict[str, tuple[list[str], list[list]]]) -> bytes:
    wb = Workbook()
    wb.remove(wb.active)
    for title, (headers, rows) in sheets.items():
        ws = wb.create_sheet(title[:31])
        ws.append(headers)
        for c in ws[1]:
            c.font = Font(bold=True, color="FFFFFF", size=10)
            c.fill = PatternFill("solid", fgColor=NAVY)
            c.alignment = Alignment(vertical="center", wrap_text=True)
        for r in rows:
            ws.append(r)
        widths = [len(str(h)) for h in headers]
        for r in rows[:400]:
            for i, v in enumerate(r):
                if i < len(widths):
                    widths[i] = max(widths[i], min(60, len(str(v))))
        for i, wdt in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(i)].width = min(62, wdt + 3)
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


# -------------------------------------------------------------------- PDF
def to_pdf(title: str, subtitle: str, blocks: list) -> bytes:
    """blocks: list of ("heading", str) | ("para", str) | ("table", headers, rows)"""
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=landscape(A4),
                            leftMargin=14 * mm, rightMargin=14 * mm,
                            topMargin=14 * mm, bottomMargin=14 * mm,
                            title=title, author="Gsure Technologies Private Limited")
    ss = getSampleStyleSheet()
    h1 = ParagraphStyle("h1", parent=ss["Heading1"], fontName="Helvetica-Bold",
                        fontSize=16, textColor=colors.HexColor("#0E2A47"),
                        spaceAfter=2)
    sub = ParagraphStyle("sub", parent=ss["Normal"], fontSize=9,
                         textColor=colors.HexColor("#5C6672"), spaceAfter=10)
    h2 = ParagraphStyle("h2", parent=ss["Heading2"], fontName="Helvetica-Bold",
                        fontSize=11, textColor=colors.HexColor("#1E5C9E"),
                        spaceBefore=10, spaceAfter=4)
    body = ParagraphStyle("body", parent=ss["Normal"], fontSize=8.5, leading=11)

    story = [Paragraph(title, h1), Paragraph(subtitle, sub)]
    for block in blocks:
        if block[0] == "heading":
            story.append(Paragraph(block[1], h2))
        elif block[0] == "para":
            story.append(Paragraph(block[1], body))
            story.append(Spacer(1, 4))
        elif block[0] == "table":
            _, headers, rows = block
            data = [[Paragraph(f"<b>{h}</b>", ParagraphStyle(
                "th", parent=body, textColor=colors.white, fontSize=8))
                for h in headers]]
            for r in rows:
                data.append([Paragraph(str(v), body) for v in r])
            t = Table(data, repeatRows=1, hAlign="LEFT")
            t.setStyle(TableStyle([
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#0E2A47")),
                ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#CBD8E6")),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1),
                 [colors.white, colors.HexColor("#F1F5FA")]),
                ("LEFTPADDING", (0, 0), (-1, -1), 5),
                ("RIGHTPADDING", (0, 0), (-1, -1), 5),
                ("TOPPADDING", (0, 0), (-1, -1), 3),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
            ]))
            story.append(t)
            story.append(Spacer(1, 8))

    def footer(canvas, _doc):
        canvas.saveState()
        canvas.setFont("Helvetica", 7)
        canvas.setFillColor(colors.HexColor("#5C6672"))
        canvas.drawString(14 * mm, 8 * mm,
                          "Confidential  |  Gsure Technologies Private Limited")
        canvas.drawRightString(landscape(A4)[0] - 14 * mm, 8 * mm,
                               f"Page {canvas.getPageNumber()}")
        canvas.restoreState()

    doc.build(story, onFirstPage=footer, onLaterPages=footer)
    return buf.getvalue()


# --------------------------------------------------------- report builders
METRIC_LABELS = [
    ("targets", "Target groups"),
    ("M1_groups_reached", "M1 Groups reached"),
    ("M1_reach_pct", "M1 Reach %"),
    ("M2_failures", "M2 Delivery failures"),
    ("M2_failure_pct", "M2 Failure %"),
    ("M3_pending", "M3 Pending deliveries"),
    ("M9_groups_read", "M9 Groups with read receipt"),
    ("M9_read_pct_of_reached", "M9 Read % of reached"),
    ("M9_member_reads", "M9 Member level reads (aggregate)"),
    ("M4_clicks", "M4 Link clicks"),
    ("M4_unique_clicks", "M4 Unique clicks"),
    ("M4_groups_clicked", "M4 Groups that clicked"),
    ("M4_ctr_pct", "M4 Click through %"),
    ("M5_media_views", "M5 Tracked media views"),
    ("M6_replies", "M6 Replies"),
    ("M6_optouts", "M6 Opt out requests"),
    ("M7_engagement_index", "M7 Engagement index"),
]


def campaign_report_data(campaign_id: int) -> dict:
    m = analytics.campaign_metrics(campaign_id)
    snaps = analytics.snapshots_for(campaign_id)
    perf = analytics.group_performance(campaign_id, limit=5000)
    return {"metrics": m, "snapshots": snaps, "groups": perf}


def campaign_sheets(campaign_id: int) -> dict:
    d = campaign_report_data(campaign_id)
    m = d["metrics"]
    summary_rows = [[label, m.get(key, "")] for key, label in METRIC_LABELS]
    snap_headers = ["Interval", "Taken at", "Reached", "Failed", "Pending",
                    "Groups read", "Clicks", "Media views", "Replies",
                    "Engagement index"]
    snap_rows = [[s.get("label"), s.get("taken_at", "")[:19],
                  s.get("M1_groups_reached", 0), s.get("M2_failures", 0),
                  s.get("M3_pending", 0), s.get("M9_groups_read", 0),
                  s.get("M4_clicks", 0), s.get("M5_media_views", 0),
                  s.get("M6_replies", 0), s.get("M7_engagement_index", 0)]
                 for s in d["snapshots"]]
    grp_headers = ["Group", "Category", "Members", "Status", "Attempts",
                   "Member reads", "Clicks", "Media views", "Interactions",
                   "Error code", "Error"]
    grp_rows = [[g["group"], g["category"], g["members"], g["status"],
                 g["attempts"], g["member_reads"], g["clicks"],
                 g["media_views"], g["interactions"], g["error_code"] or "",
                 g["error"] or ""] for g in d["groups"]]
    fail_headers = ["Error code", "Error title", "Groups affected"]
    fail_rows = [[f["code"], f["title"], f["count"]]
                 for f in m.get("M2_failure_breakdown", [])]
    return {
        "Summary": (["Metric", "Value"], summary_rows),
        "Interval snapshots": (snap_headers, snap_rows),
        "Per group log": (grp_headers, grp_rows),
        "Failure analysis": (fail_headers, fail_rows),
    }


def campaign_pdf(campaign_id: int) -> bytes:
    d = campaign_report_data(campaign_id)
    m = d["metrics"]
    sheets = campaign_sheets(campaign_id)
    blocks = [
        ("heading", "Campaign summary"),
        ("table", ["Metric", "Value"],
         [[l, m.get(k, "")] for k, l in METRIC_LABELS]),
        ("heading", "Interval snapshots (15m, 1h, 3h, 8h)"),
        ("table", *sheets["Interval snapshots"]),
    ]
    if sheets["Failure analysis"][1]:
        blocks += [("heading", "Failure analysis"),
                   ("table", *sheets["Failure analysis"])]
    grp_h, grp_r = sheets["Per group log"]
    blocks += [("heading", f"Per group delivery log (first 120 of {len(grp_r)})"),
               ("table", grp_h, grp_r[:120])]
    blocks += [("para",
                "Metric identifiers M1 to M9 correspond to the feasibility "
                "table in the project proposal. M5 is measurable only because "
                "media is delivered through a tracked landing endpoint; media "
                "sent as a direct WhatsApp attachment cannot be measured.")]
    return to_pdf(
        f"Campaign report: {m.get('campaign_name', '')}",
        f"{m.get('campaign_code','')}  |  channel {m.get('channel','')}  |  "
        f"generated {dt.datetime.now():%d %b %Y %H:%M}",
        blocks)


def groups_sheets() -> dict:
    with SessionLocal() as s:
        rows = list(s.scalars(select(Group)))
    headers = ["Group id", "WhatsApp group id", "Name", "Category", "Folder",
               "Tags", "Members", "Region", "Active", "Source"]
    data = [[g.id, g.wa_group_id, g.name, g.category, g.folder,
             ", ".join(g.tags), g.member_count, g.region,
             "Yes" if g.active else "No", g.source] for g in rows]
    return {"Groups": (headers, data)}
